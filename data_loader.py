# -*- coding: utf-8 -*-
"""
data_loader.py
==============
读取赛事方提供的工作簿（Track 3 Data.xlsx），把各工作表加载为
Python 数据结构，供研判引擎使用。

设计原则
--------
1. 只读不改：原始工作簿永远不被写入（赛事规则 + 数据质量要求）。
2. 每个数据字段都保留其「来源编号」（SIM-BASE / SIM-SCALE / SIM-WEATHER 等），
   保证引擎输出的每一条判断都可以回溯到数据来源（可追溯性，评分要点之一）。
3. 加载时做基本质量检查（缺失标记、字段口径），检查结果交给引擎降级处理。
"""

import openpyxl
from dataclasses import dataclass, field
from typing import Optional


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

@dataclass
class BaselineRow:
    """01_机场月度运行基线 的一行记录。"""
    year: int
    month: int
    airport: str
    airport_class: str
    plan_ratio: Optional[float]       # 按计划完成航班的比例
    release_ratio: Optional[float]    # 机场放行正常比例（口径不同，不可混用）
    peer_avg: Optional[float]         # 同类机场平均比例
    self_avg_12m: Optional[float]     # 该机场过去12个月平均比例
    diff_pp: Optional[float]          # 与自身平均差（百分点）
    source: str                       # 来源编号，如 SIM-BASE


@dataclass
class AirportScale:
    """02_机场年度规模 的一行记录。"""
    year: int
    airport: str
    city: str
    province: str
    airport_class: str
    passengers_wan: Optional[int]     # 旅客吞吐量（万人次）
    movements: Optional[int]          # 飞机起降架次
    importance: str                   # 航空公司业务重要程度：高/中/低
    source: str


@dataclass
class WeatherRow:
    """03_机场天气月度 的一行记录。"""
    year: int
    month: int
    airport: str
    station_id: str
    precipitation_mm: Optional[float]
    temp_avg_c: Optional[float]
    wind_max_ms: Optional[float]
    adverse_days: Optional[int]
    is_missing: bool                  # 是否缺少数据（DQ002：模型必须降级）
    source: str


@dataclass
class Event:
    """04_模拟事件 的一条事件。"""
    event_id: str
    name: str
    time: str
    airport: str
    weather_desc: str
    ops_desc: str
    importance: str
    field_note: str
    known: str
    unknown: str
    status: str
    tag: str


@dataclass
class Rule:
    """05_运营规则 的一条规则。规则是引擎的硬约束，优先级高于任何评分。"""
    rule_id: str
    topic: str
    trigger: str
    checks: str
    action: str
    escalate_to: str
    human_approval: bool              # 必须人工审批
    forbidden: str                    # 禁止自动执行事项
    source: str
    status: str


@dataclass
class Feedback:
    """07_人工反馈 的一条记录，用于反馈闭环（RULE006）。"""
    event_id: str
    rec_id: str
    result: str                       # 接受 / 修改 / 拒绝
    reason: str
    final_action: str
    outcome: str
    improvement: str
    reviewer: str
    review_date: str


@dataclass
class TestQuestion:
    """08_测试问题 的一道题。"""
    test_id: str
    qtype: str
    question: str
    expected: str
    forbidden: str
    fact_check: str
    score_point: str


# ---------------------------------------------------------------------------
# 加载器
# ---------------------------------------------------------------------------

class DataLoader:
    """一次性加载工作簿内全部相关表格。"""

    SHEET_BASELINE = '01_机场月度运行基线'
    SHEET_SCALE = '02_机场年度规模'
    SHEET_WEATHER = '03_机场天气月度'
    SHEET_EVENT = '04_模拟事件'
    SHEET_RULE = '05_运营规则'
    SHEET_FEEDBACK = '07_人工反馈'
    SHEET_TEST = '08_测试问题'

    def __init__(self, path: str):
        self.path = path
        wb = openpyxl.load_workbook(path, data_only=True)
        self.baselines = self._load_baseline(wb)
        self.scales = self._load_scale(wb)
        self.weather = self._load_weather(wb)
        self.events = self._load_events(wb)
        self.rules = self._load_rules(wb)
        self.feedback = self._load_feedback(wb)
        self.tests = self._load_tests(wb)

    # -- 各表加载 -----------------------------------------------------------

    def _load_baseline(self, wb):
        out = []
        for r in wb[self.SHEET_BASELINE].iter_rows(min_row=2, values_only=True):
            if r[2] is None:
                continue
            out.append(BaselineRow(
                year=int(r[0]), month=int(r[1]), airport=r[2], airport_class=r[3],
                plan_ratio=r[4], release_ratio=r[5], peer_avg=r[6],
                self_avg_12m=r[7], diff_pp=r[8], source=r[10] or 'SIM-BASE',
            ))
        return out

    def _load_scale(self, wb):
        out = []
        for r in wb[self.SHEET_SCALE].iter_rows(min_row=2, values_only=True):
            if r[1] is None:
                continue
            out.append(AirportScale(
                year=int(r[0]), airport=r[1], city=r[2], province=r[3],
                airport_class=r[4], passengers_wan=r[5], movements=r[6],
                importance=r[7], source=r[8] or 'SIM-SCALE',
            ))
        return out

    def _load_weather(self, wb):
        out = []
        for r in wb[self.SHEET_WEATHER].iter_rows(min_row=2, values_only=True):
            if r[2] is None:
                continue
            out.append(WeatherRow(
                year=int(r[0]), month=int(r[1]), airport=r[2], station_id=r[3],
                precipitation_mm=r[4], temp_avg_c=r[5], wind_max_ms=r[6],
                adverse_days=r[7], is_missing=(r[8] == '是'), source=r[9] or 'SIM-WEATHER',
            ))
        return out

    def _load_events(self, wb):
        out = []
        for r in wb[self.SHEET_EVENT].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            out.append(Event(
                event_id=r[0], name=r[1], time=r[2], airport=r[3],
                weather_desc=r[4] or '', ops_desc=r[5] or '', importance=r[6],
                field_note=r[7] or '', known=r[8] or '', unknown=r[9] or '',
                status=r[10], tag=r[11],
            ))
        return out

    def _load_rules(self, wb):
        out = []
        for r in wb[self.SHEET_RULE].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            out.append(Rule(
                rule_id=r[0], topic=r[1], trigger=r[2], checks=r[3],
                action=r[4], escalate_to=r[5],
                human_approval=(r[6] == '是'), forbidden=r[7] or '',
                source=r[8] or '赛事规则', status=r[9],
            ))
        return out

    def _load_feedback(self, wb):
        out = []
        for r in wb[self.SHEET_FEEDBACK].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            out.append(Feedback(
                event_id=r[0], rec_id=r[1], result=r[2], reason=r[3] or '',
                final_action=r[4] or '', outcome=r[5] or '', improvement=r[6] or '',
                reviewer=r[7] or '', review_date=r[8] or '',
            ))
        return out

    def _load_tests(self, wb):
        out = []
        for r in wb[self.SHEET_TEST].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            out.append(TestQuestion(
                test_id=r[0], qtype=r[1], question=r[2], expected=r[3],
                forbidden=r[4], fact_check=r[5], score_point=r[6],
            ))
        return out

    # -- 查询接口 -----------------------------------------------------------

    def baseline_rows(self, airport: str):
        return sorted([b for b in self.baselines if b.airport == airport],
                      key=lambda b: (b.year, b.month))

    def latest_scale(self, airport: str) -> Optional[AirportScale]:
        rows = [s for s in self.scales if s.airport == airport]
        return max(rows, key=lambda s: s.year) if rows else None

    def weather_rows(self, airport: str):
        return sorted([w for w in self.weather if w.airport == airport],
                      key=lambda w: (w.year, w.month))

    def event(self, event_id: str) -> Optional[Event]:
        for e in self.events:
            if e.event_id == event_id:
                return e
        return None

    def rule(self, rule_id: str) -> Optional[Rule]:
        for r in self.rules:
            if r.rule_id == rule_id:
                return r
        return None
